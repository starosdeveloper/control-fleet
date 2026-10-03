"""
GPS-импорт показаний из Excel (ТЗ §8): чтение .xlsx, сопоставление колонок,
предпросмотр и применение партии. Доступы к БД и к файлу — здесь; классификация
строк — в чистом `gpsimport`.

Предпросмотр не меняет данные (ТЗ §8). Применение атомарно: при технической
ошибке записи партия откатывается целиком. Повторная загрузка того же файла не
создаёт дублей — полный дубль VIN+дата+пробег+источник отсекается.

Импорт НЕ создаёт фото, НЕ подтверждает ТО и НЕ отмечает месячный отчёт
выполненным. Показания водителя и GPS не перезаписывают друг друга — это разные
записи с разным источником.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date

from django.conf import settings
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from .gpsimport import CellValue, RowCategory, classify_import_row, normalize_date, normalize_mileage
from .models import ImportBatch, OdometerReading, Vehicle
from .storage import open_abs, save_import_file  # noqa: F401 (реэкспорт для вызовов импорта)

MAX_ROWS = 10_000
MAX_BYTES = 10 * 1024 * 1024

# Заголовки по умолчанию; администратор может указать свои при загрузке сырого файла.
DEFAULT_MAP = {"vin": "VIN", "odometer": "Пробег одометр", "date": "Дата показания"}


@dataclass
class PreviewRow:
    excel_row: int            # номер строки в файле (с учётом заголовка)
    vin: str
    mileage_display: str
    date_display: str
    category: RowCategory
    reason: str


@dataclass
class ImportResult:
    ready: int = 0
    review: int = 0
    duplicate: int = 0
    error: int = 0
    total: int = 0
    rows: list[PreviewRow] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)   # для Excel ошибок
    fatal: str = ""                                     # непусто → файл не прочитан

    def counts(self) -> dict:
        return {
            "total": self.total,
            "ready": self.ready,
            "review": self.review,
            "duplicate": self.duplicate,
            "error": self.error,
        }


# --- Чтение файла -------------------------------------------------------------

def _match_columns(header: list, mapping: dict) -> dict | None:
    """Сопоставить логические поля с индексами колонок по названиям заголовка."""
    norm = {
        (str(h).strip().lower() if h is not None else ""): idx
        for idx, h in enumerate(header)
    }
    cols = {}
    for field_name in ("vin", "odometer", "date"):
        want = (mapping.get(field_name) or "").strip().lower()
        if field_name == "date" and not want:
            cols["date"] = None  # колонки даты нет — используем общую дату
            continue
        if want in norm:
            cols[field_name] = norm[want]
        else:
            return None
    return cols


def _cell(row: tuple, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def build_preview(
    rel_path: str,
    mapping: dict,
    common_date: date | None,
    today: date,
) -> ImportResult:
    """
    Прочитать файл и классифицировать строки, ничего не записывая (ТЗ §8).
    `rel_path` — путь в закрытом хранилище к ранее сохранённому исходному файлу.
    """
    result = ImportResult()
    path = open_abs(rel_path)
    if not path.exists():
        result.fatal = "Файл импорта не найден."
        return result

    try:
        # data_only=False: формулы НЕ вычисляем — они приходят строкой «=…» и
        # помечаются ошибкой нормализацией (ТЗ §8).
        wb = load_workbook(path, read_only=True, data_only=False)
    except Exception:
        result.fatal = "Не удалось прочитать .xlsx (повреждён или не формат)."
        return result

    ws = wb.worksheets[0]  # первый лист (ТЗ §8)
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = list(next(rows_iter))
    except StopIteration:
        wb.close()
        result.fatal = "Пустой файл: нет строки заголовков."
        return result

    cols = _match_columns(header, mapping)
    if cols is None:
        wb.close()
        result.fatal = (
            "Не найдены колонки. Нужны: "
            f"VIN=«{mapping.get('vin')}», пробег=«{mapping.get('odometer')}»"
            + (f", дата=«{mapping.get('date')}»" if mapping.get("date") else "")
        )
        return result

    # Известные VIN и их существующие записи — одним заходом в БД.
    known = {v.vin: v for v in Vehicle.objects.all()}
    gps_sigs: dict[str, set] = {}
    accepted_by_vin: dict[str, list] = {}
    for r in OdometerReading.objects.exclude(
        import_batch__status=ImportBatch.Status.REVERTED
    ).values_list("vehicle__vin", "measured_at", "value_km", "source", "verification"):
        vin, mdate, val, source, verif = r
        if source == OdometerReading.Source.GPS:
            gps_sigs.setdefault(vin, set()).add((mdate, val))
        if verif == OdometerReading.Verification.ACCEPTED:
            accepted_by_vin.setdefault(vin, []).append((val, mdate))

    # Накопители для дублей внутри самого файла (строка-близнец в одном файле).
    seen_in_file: dict[str, set] = {}

    excel_row = 1
    for raw in rows_iter:
        excel_row += 1
        if result.total >= MAX_ROWS:
            result.fatal = f"Превышен предел {MAX_ROWS} строк."
            break
        # Полностью пустая строка в конце файла — пропускаем молча.
        if raw is None or all(c is None or (isinstance(c, str) and not c.strip()) for c in raw):
            continue
        result.total += 1

        raw_vin = _cell(raw, cols["vin"])
        vin = str(raw_vin).strip().upper() if raw_vin is not None else ""
        mileage = normalize_mileage(_cell(raw, cols["odometer"]))
        when = normalize_date(_cell(raw, cols.get("date")), common_date, today)

        vin_known = vin in known
        sigs = set(gps_sigs.get(vin, set())) | seen_in_file.get(vin, set())
        same_date_vals = []
        prior = accepted_by_vin.get(vin, [])
        if when.value is not None:
            same_date_vals = [v for (v, d) in prior if d == when.value]

        category, reason = classify_import_row(
            known_vin=vin_known,
            mileage=mileage,
            when=when,
            gps_signatures=sigs,
            same_date_accepted_values=same_date_vals,
            prior_accepted=prior,
            jump_limit=settings.ODOMETER_DAILY_JUMP_KM,
            conflict_limit=settings.GPS_DRIVER_CONFLICT_KM,
        )

        pr = PreviewRow(
            excel_row=excel_row,
            vin=vin or "—",
            mileage_display=(str(mileage.value) if mileage.value is not None else str(_cell(raw, cols["odometer"]))),
            date_display=(when.value.strftime("%d.%m.%Y") if when.value else "—"),
            category=category,
            reason=reason,
        )
        result.rows.append(pr)

        if category == RowCategory.ERROR:
            result.error += 1
            result.errors.append(
                {"row": excel_row, "vin": vin, "value": pr.mileage_display, "reason": reason}
            )
        elif category == RowCategory.DUPLICATE:
            result.duplicate += 1
        elif category == RowCategory.REVIEW:
            result.review += 1
        else:
            result.ready += 1
            # Готовую строку учитываем как уже «виденную» в файле, чтобы её близнец
            # ниже стал дублем, а не второй принятой записью.
            if when.value is not None and mileage.value is not None:
                seen_in_file.setdefault(vin, set()).add((when.value, mileage.value))
                accepted_by_vin.setdefault(vin, []).append((mileage.value, when.value))

    wb.close()
    return result


# --- Применение ---------------------------------------------------------------

@transaction.atomic
def apply_import(
    rel_path: str,
    mapping: dict,
    common_date: date | None,
    author,
    today: date,
) -> tuple[ImportBatch, ImportResult]:
    """
    Применить импорт (ТЗ §8). Пересчитывает предпросмотр в транзакции (источник
    истины — тот же файл) и пишет: READY — принятые, REVIEW — на проверке, дубли
    и ошибки пропускает. Партия хранит исходный файл, автора, дату, итог. При
    технической ошибке вся транзакция откатывается.
    """
    result = build_preview(rel_path, mapping, common_date, today)
    if result.fatal:
        raise ValueError(result.fatal)

    batch = ImportBatch.objects.create(
        kind=ImportBatch.Kind.GPS,
        source_path=rel_path,
        author=author,
        status=ImportBatch.Status.ACTIVE,
        summary={"counts": result.counts(), "errors": result.errors},
    )

    known = {v.vin: v for v in Vehicle.objects.all()}
    to_create = []
    now = timezone.now()
    for pr in result.rows:
        if pr.category not in (RowCategory.READY, RowCategory.REVIEW):
            continue
        vehicle = known.get(pr.vin)
        if vehicle is None:
            continue
        # date_display → date; mileage_display → int (оба уже нормализованы).
        day = _parse_display_date(pr.date_display)
        value = int(pr.mileage_display)
        verification = (
            OdometerReading.Verification.ACCEPTED
            if pr.category == RowCategory.READY
            else OdometerReading.Verification.PENDING
        )
        to_create.append(
            OdometerReading(
                vehicle=vehicle,
                source=OdometerReading.Source.GPS,
                value_km=value,
                measured_at=day,
                received_at=now,
                import_batch=batch,
                verification=verification,
                note=pr.reason if pr.category == RowCategory.REVIEW else "",
            )
        )

    OdometerReading.objects.bulk_create(to_create)
    return batch, result


def _parse_display_date(s: str) -> date:
    from datetime import datetime

    return datetime.strptime(s, "%d.%m.%Y").date()


@transaction.atomic
def revert_batch(batch: ImportBatch, staff, reason: str) -> ImportBatch:
    """
    Отменить партию (ТЗ §8): записи исключаются из расчёта, но остаются в истории.
    Актуальный пробег и ТО пересчитываются автоматически, т.к. выборка принятых
    показаний исключает отменённые партии (`Vehicle.accepted_reading_rows`).

    Откат определён только для GPS-партий: это переключение видимости показаний,
    а не удаление сущностей. У партии подключения парка (`Kind.FLEET`) откат создал
    бы ТС/назначения/менеджеров — такую отмену ТЗ не описывает, и тихо развязывать
    назначения через общий механизм небезопасно. Отменить её результат можно только
    вручную через карточку ТС (закрыть назначение, архивировать ТС) — это
    осознанное ограничение, не пробел.
    """
    if batch.kind != ImportBatch.Kind.GPS:
        raise ValueError(
            "Откат доступен только для партий GPS-показаний. Подключение парка "
            "отменяется вручную через карточку ТС."
        )
    locked = ImportBatch.objects.select_for_update().get(pk=batch.pk)
    if locked.status == ImportBatch.Status.REVERTED:
        return locked
    locked.status = ImportBatch.Status.REVERTED
    locked.reverted_by = staff
    locked.reverted_at = timezone.now()
    locked.revert_reason = reason
    locked.save(update_fields=["status", "reverted_by", "reverted_at", "revert_reason"])
    return locked


# --- Шаблон и Excel ошибок ----------------------------------------------------

def build_template() -> bytes:
    """Шаблон .xlsx с нужными заголовками и строкой-примером (ТЗ §8)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "GPS"
    ws.append([DEFAULT_MAP["vin"], DEFAULT_MAP["odometer"], DEFAULT_MAP["date"]])
    ws.append(["TRUCKFLEET0000001", 170000, "2026-09-30"])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def errors_xlsx(errors: list[dict]) -> bytes:
    """
    Excel ошибок с номером исходной строки и причиной (ТЗ §8). Общий для обоих
    импортов (GPS и подключение парка — §7); `field` у GPS-ошибок пустой, у
    подключения парка называет проблемную колонку.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Ошибки"
    ws.append(["№ строки", "VIN", "Поле", "Значение", "Причина"])
    for e in errors:
        ws.append([e.get("row"), e.get("vin"), e.get("field", ""), e.get("value"), e.get("reason")])
    for col, width in zip("ABCDE", (10, 22, 16, 14, 44)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
