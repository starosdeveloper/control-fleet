"""
Массовое подключение парка из Excel (ТЗ §7): первичный импорт VIN, ИНН,
менеджера, госномера, модели, пробега и даты показания, даты и пробега ТО.

Доступы к БД и к файлу — здесь; классификация строк и разбор пар значений — в
чистом `onboarding`. Предпросмотр не меняет данные; применение атомарно —
при ошибке партия откатывается целиком (ТЗ §8, тот же принцип применён и
здесь). Повторная загрузка того же файла безопасна: READY становится EXISTS
на уже подключённые ТС, CONFLICT не трогает чужие назначения.

Созданная строка показаний получает `source=manual` (это ввод сотрудника, не
отчёт водителя и не GPS) и `verification=accepted` без прогона через
`readings.classify_reading` — это стартовая база нового ТС, конфликтовать не с
чем, тот же принцип, что у `seed_demo`. Созданное ТО — сразу `status=confirmed`,
`basis=initial` — так же, как «начальная база» в демо-данных.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime

from django.contrib.auth import get_user_model
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from .access import MANAGER_GROUP
from .gpsservice import MAX_BYTES, MAX_ROWS  # те же лимиты, что у GPS-импорта (§8)
from .models import Assignment, ImportBatch, ManagerAssignment, Maintenance, OdometerReading, Tenant, Vehicle
from .onboarding import (
    RowCategory,
    classify_onboarding_row,
    cross_check_maint_vs_reading,
    normalize_optional_pair,
)
from .storage import open_abs, save_import_file  # noqa: F401 (реэкспорт, как в gpsservice)
from .validators import is_valid_inn, is_valid_vin, normalize_inn, normalize_vin

DEFAULT_MAP = {
    "vin": "VIN",
    "inn": "ИНН",
    "manager": "Менеджер",
    "plate": "Госномер",
    "model": "Модель",
    "mileage": "Пробег",
    "mileage_date": "Дата показания",
    "maint_date": "Дата ТО",
    "maint_mileage": "Пробег ТО",
}
_REQUIRED = ("vin", "inn", "manager")


@dataclass
class PreviewRow:
    excel_row: int
    vin: str
    inn: str
    manager: str
    plate: str
    model: str
    mileage: int | None
    mileage_date: date | None
    maint_date: date | None
    maint_mileage: int | None
    category: RowCategory
    reason: str


@dataclass
class ImportResult:
    ready: int = 0
    exists: int = 0
    conflict: int = 0
    error: int = 0
    total: int = 0
    rows: list[PreviewRow] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)
    fatal: str = ""

    def counts(self) -> dict:
        return {
            "total": self.total, "ready": self.ready, "exists": self.exists,
            "conflict": self.conflict, "error": self.error,
        }


def _cell(row: tuple, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _match_columns(header: list, mapping: dict) -> dict | None:
    norm = {
        (str(h).strip().lower() if h is not None else ""): idx
        for idx, h in enumerate(header)
    }
    cols: dict[str, int | None] = {}
    for key, want in mapping.items():
        want_norm = (want or "").strip().lower()
        if not want_norm:
            cols[key] = None
            continue
        if want_norm in norm:
            cols[key] = norm[want_norm]
        elif key in _REQUIRED:
            return None
        else:
            cols[key] = None
    return cols


def build_preview(rel_path: str, mapping: dict, today: date) -> ImportResult:
    """Прочитать файл и классифицировать строки, ничего не записывая (ТЗ §7, §8)."""
    result = ImportResult()
    path = open_abs(rel_path)
    if not path.exists():
        result.fatal = "Файл импорта не найден."
        return result

    try:
        wb = load_workbook(path, read_only=True, data_only=False)
    except Exception:
        result.fatal = "Не удалось прочитать .xlsx (повреждён или не формат)."
        return result

    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = list(next(rows_iter))
    except StopIteration:
        wb.close()
        result.fatal = "Пустой файл: нет строки заголовков."
        return result

    cols = _match_columns(header, mapping)
    if cols is None:
        wb.close()
        result.fatal = (
            "Не найдены обязательные колонки. Нужны: "
            f"VIN=«{mapping.get('vin')}», ИНН=«{mapping.get('inn')}», "
            f"менеджер=«{mapping.get('manager')}»."
        )
        return result

    active_tenant_by_vin = {
        a.vehicle.vin: a.tenant.inn
        for a in Assignment.objects.filter(ended_at__isnull=True).select_related("vehicle", "tenant")
    }
    manager_users = set(
        get_user_model().objects.filter(groups__name=MANAGER_GROUP).values_list("username", flat=True)
    )

    seen_vins: set[str] = set()
    excel_row = 1
    for raw in rows_iter:
        excel_row += 1
        if result.total >= MAX_ROWS:
            result.fatal = f"Превышен предел {MAX_ROWS} строк."
            break
        if raw is None or all(c is None or (isinstance(c, str) and not c.strip()) for c in raw):
            continue
        result.total += 1

        vin_raw = _cell(raw, cols["vin"])
        vin = normalize_vin(str(vin_raw)) if vin_raw is not None else ""
        vin_ok = is_valid_vin(vin)

        inn_raw = _cell(raw, cols["inn"])
        inn = normalize_inn(str(inn_raw)) if inn_raw is not None else ""
        inn_ok = is_valid_inn(inn)

        manager_raw = _cell(raw, cols["manager"])
        manager = str(manager_raw).strip() if manager_raw is not None else ""
        manager_known = manager in manager_users

        plate_raw = _cell(raw, cols.get("plate"))
        plate = str(plate_raw).strip() if plate_raw is not None else ""
        model_raw = _cell(raw, cols.get("model"))
        model = str(model_raw).strip() if model_raw is not None else ""

        reading = normalize_optional_pair(
            _cell(raw, cols.get("mileage")), _cell(raw, cols.get("mileage_date")), today, "пробег/дата показания"
        )
        maint = normalize_optional_pair(
            _cell(raw, cols.get("maint_mileage")), _cell(raw, cols.get("maint_date")), today, "дата/пробег ТО"
        )
        cross_error = cross_check_maint_vs_reading(reading.mileage, maint.mileage)

        duplicate_in_file = vin_ok and vin in seen_vins
        if vin_ok:
            seen_vins.add(vin)

        category, reason = classify_onboarding_row(
            vin_ok=vin_ok,
            inn_ok=inn_ok,
            manager_known=manager_known,
            reading_error=reading.error,
            maint_error=maint.error,
            cross_error=cross_error,
            duplicate_in_file=duplicate_in_file,
            existing_active_tenant_inn=active_tenant_by_vin.get(vin) if vin_ok else None,
            row_inn=inn,
        )

        pr = PreviewRow(
            excel_row=excel_row, vin=vin or "—", inn=inn or "—", manager=manager or "—",
            plate=plate, model=model, mileage=reading.mileage, mileage_date=reading.when,
            maint_date=maint.when, maint_mileage=maint.mileage, category=category, reason=reason,
        )
        result.rows.append(pr)

        if category == RowCategory.ERROR:
            result.error += 1
            if not vin_ok:
                field_name = "VIN"
            elif not inn_ok:
                field_name = "ИНН"
            elif not manager_known:
                field_name = "менеджер"
            elif reading.error:
                field_name = "пробег/дата показания"
            elif maint.error:
                field_name = "дата/пробег ТО"
            elif cross_error:
                field_name = "пробег ТО"
            else:
                field_name = "VIN"  # duplicate_in_file
            result.errors.append(
                {"row": excel_row, "vin": vin, "field": field_name, "value": "", "reason": reason}
            )
        elif category == RowCategory.CONFLICT:
            result.conflict += 1
        elif category == RowCategory.EXISTS:
            result.exists += 1
        else:
            result.ready += 1

    wb.close()
    return result


@transaction.atomic
def apply_import(rel_path: str, mapping: dict, author, today: date) -> tuple[ImportBatch, ImportResult]:
    """
    Применить подключение парка (ТЗ §7). READY-строки создают/переоткрывают ТС,
    арендатора (get_or_create по ИНН), назначение и закрепление менеджера, и, если
    указаны, начальные показание и ТО. EXISTS и CONFLICT не меняют ничего — это
    отчётные категории предпросмотра. При технической ошибке откат целиком.
    """
    result = build_preview(rel_path, mapping, today)
    if result.fatal:
        raise ValueError(result.fatal)

    batch = ImportBatch.objects.create(
        kind=ImportBatch.Kind.FLEET,
        source_path=rel_path,
        author=author,
        status=ImportBatch.Status.ACTIVE,
        summary={"counts": result.counts(), "errors": result.errors},
    )

    managers_by_name = {
        u.username: u
        for u in get_user_model().objects.filter(groups__name=MANAGER_GROUP)
    }
    now = timezone.now()

    for pr in result.rows:
        if pr.category != RowCategory.READY:
            continue

        vehicle, _ = Vehicle.objects.get_or_create(
            vin=pr.vin, defaults={"plate": pr.plate, "model": pr.model}
        )
        tenant, _ = Tenant.objects.get_or_create(inn=pr.inn)
        assignment = Assignment.objects.create(vehicle=vehicle, tenant=tenant, started_at=now)
        ManagerAssignment.objects.create(
            assignment=assignment, manager=managers_by_name[pr.manager], started_at=now
        )

        if pr.mileage is not None and pr.mileage_date is not None:
            OdometerReading.objects.create(
                vehicle=vehicle, source=OdometerReading.Source.MANUAL, value_km=pr.mileage,
                measured_at=pr.mileage_date, received_at=now, import_batch=batch,
                verification=OdometerReading.Verification.ACCEPTED,
                note="начальный пробег при подключении парка",
            )
        if pr.maint_mileage is not None and pr.maint_date is not None:
            Maintenance.objects.create(
                vehicle=vehicle, performed_on=pr.maint_date, mileage_km=pr.maint_mileage,
                status=Maintenance.Status.CONFIRMED, basis=Maintenance.Basis.INITIAL,
                confirmed_by=author, confirmed_at=now, reason="начальная база при подключении парка",
            )

    return batch, result


# --- Шаблон и Excel ошибок ----------------------------------------------------

def build_template() -> bytes:
    """Шаблон .xlsx с колонками §7 и строкой-примером."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Парк"
    ws.append(list(DEFAULT_MAP.values()))
    ws.append([
        "TRUCKFLEET0000011", "7700000099", "manager1", "А011АА750", "КАМАЗ-5490",
        140000, "2026-10-01", "2026-06-15", 120000,
    ])
    for col, width in zip(
        "ABCDEFGHI", (20, 14, 14, 12, 16, 10, 16, 14, 12)
    ):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
