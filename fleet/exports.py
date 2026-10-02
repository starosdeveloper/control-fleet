"""
Экспорт для кабинета сотрудника (ТЗ §7): отфильтрованная таблица в Excel и
выбранный отчёт с вложениями одним ZIP.

Excel строится из уже собранных строк дашборда (`dashboard.DashboardRow`) — те же
значения, что видит сотрудник на экране, тем же срезом. ZIP берёт оригиналы
вложений из закрытого хранилища; имена внутри архива осмысленные, но файлы те же.
"""
from __future__ import annotations

import io
import zipfile
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font

from .dashboard import DashboardRow
from .maintenance import MaintStatus
from .models import Attachment, Report
from .storage import open_abs

# Заголовки колонок таблицы контроля (ТЗ §7), в том же порядке, что на экране.
TABLE_HEADERS = [
    "ID",
    "VIN",
    "Госномер",
    "ИНН",
    "Менеджер",
    "Последний отчёт",
    "Пробег, км",
    "Дата пробега",
    "Источник",
    "Дата ТО",
    "Пробег ТО, км",
    "Следующее ТО, км",
    "Остаток, км",
    "Статус ТО",
    "Отчёт в месяце",
    "Фото в месяце",
    "Комментарий",
]

_SOURCE_RU = {"driver": "водитель", "gps": "GPS", "manual": "ручной", "": "—"}


def _fmt_date(d: date | None) -> str:
    return d.strftime("%d.%m.%Y") if d else ""


def _row_values(row: DashboardRow) -> list:
    v = row.vehicle
    return [
        v.id,
        v.vin,
        v.plate or "",
        row.tenant_inn or "",
        row.manager_name or "",
        _fmt_date(row.last_report_date),
        row.mileage_km if row.mileage_km is not None else "",
        _fmt_date(row.mileage_date),
        _SOURCE_RU.get(row.mileage_source, row.mileage_source or "—"),
        _fmt_date(row.maint_date),
        row.maint_mileage_km if row.maint_mileage_km is not None else "",
        row.maint.next_due_km if row.maint.next_due_km is not None else "",
        row.maint.remainder_km if row.maint.remainder_km is not None else "",
        row.maint.text,
        "да" if row.reported_in_month else "нет",
        "да" if row.photos_in_month else "нет",
        "да" if row.has_comment else "нет",
    ]


def control_table_xlsx(rows: list[DashboardRow], month_first: date) -> bytes:
    """Собрать .xlsx отфильтрованной таблицы контроля. Возвращает байты файла."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"Контроль {month_first:%Y-%m}"

    ws.append(TABLE_HEADERS)
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold

    for row in rows:
        ws.append(_row_values(row))

    # Ширины по содержимому заголовка — читаемо, без тяжёлого авто-фита.
    widths = [6, 19, 12, 14, 12, 16, 11, 13, 10, 13, 13, 16, 12, 24, 15, 15, 13]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_KIND_RU = {
    Attachment.Kind.DASHBOARD: "одометр",
    Attachment.Kind.VEHICLE: "ракурс",
    Attachment.Kind.DEFECT: "дефект",
    Attachment.Kind.WORK_ORDER: "заказ-наряд",
    Attachment.Kind.SERVICE_BOOK: "сервисная-книжка",
}
_VIEW_RU = {
    Attachment.View.FRONT: "спереди",
    Attachment.View.BACK: "сзади",
    Attachment.View.LEFT: "слева",
    Attachment.View.RIGHT: "справа",
    Attachment.View.NONE: "",
}


def _manifest_text(report: Report) -> str:
    v = report.vehicle
    lines = [
        f"Отчёт {report.number}",
        f"ТС: ID {v.id}, VIN {v.vin}" + (f", госномер {v.plate}" if v.plate else ""),
        f"Отчётный месяц: {report.period_month:%Y-%m}",
        f"Состояние: {report.get_state_display()}",
        f"Отправлен: {report.submitted_at:%d.%m.%Y %H:%M}" if report.submitted_at else "Отправлен: —",
        f"Показание одометра: {report.odometer_value_km} км"
        if report.odometer_value_km is not None
        else "Показание одометра: —",
        f"Замечания: {report.get_defects_display()}",
    ]
    if report.comment.strip():
        lines.append(f"Комментарий: {report.comment.strip()}")
    if report.maint_choice:
        lines.append(f"Сведения о ТО: {report.get_maint_choice_display()}")
        if report.maint_choice == Report.MaintChoice.DONE:
            lines.append(
                f"ТО: {report.maint_date:%d.%m.%Y}, пробег {report.maint_mileage_km} км"
                + (f", заказ-наряд {report.maint_work_order}" if report.maint_work_order else "")
            )
    return "\n".join(lines) + "\n"


def report_zip(report: Report) -> bytes:
    """
    Собрать ZIP выбранного отчёта: manifest.txt + все вложения под осмысленными
    именами. Пропавший на диске файл пропускается с отметкой в манифесте, архив не
    рушится.
    """
    buf = io.BytesIO()
    missing: list[str] = []
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        counters: dict[str, int] = {}
        for att in report.attachments.all().order_by("kind", "view", "id"):
            kind_ru = _KIND_RU.get(att.kind, att.kind)
            view_ru = _VIEW_RU.get(att.view, "")
            base = f"{kind_ru}_{view_ru}".strip("_") if view_ru else kind_ru
            counters[base] = counters.get(base, 0) + 1
            suffix = "" if counters[base] == 1 else f"_{counters[base]}"
            ext = (att.stored_path.rsplit(".", 1)[-1] or "bin")
            arcname = f"{report.number}/{base}{suffix}.{ext}"

            path = open_abs(att.stored_path)
            if not path.exists():
                missing.append(att.stored_path)
                continue
            zf.write(path, arcname)

        manifest = _manifest_text(report)
        if missing:
            manifest += "\nОтсутствуют на диске:\n" + "\n".join(missing) + "\n"
        zf.writestr(f"{report.number}/manifest.txt", manifest)

    return buf.getvalue()
